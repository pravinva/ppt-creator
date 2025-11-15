"""
Test integration of Diagrams library with PowerPoint
Shows how to embed high-quality architecture diagrams into presentations
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor


def create_presentation_with_diagrams():
    """Create a PowerPoint with the generated architecture diagrams"""
    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Databricks colors
    DATABRICKS_RED = RGBColor(255, 54, 33)
    DATABRICKS_DARK = RGBColor(27, 49, 57)

    # Slide 1: Title
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout

    # Add title
    title_box = slide1.shapes.add_textbox(Inches(0.5), Inches(2), Inches(9), Inches(1))
    title_frame = title_box.text_frame
    title_frame.text = "Enterprise Architecture Diagrams"
    title_para = title_frame.paragraphs[0]
    title_para.font.size = Pt(54)
    title_para.font.bold = True
    title_para.font.color.rgb = DATABRICKS_RED
    title_para.alignment = PP_ALIGN.CENTER

    # Add subtitle
    subtitle_box = slide1.shapes.add_textbox(Inches(0.5), Inches(3.5), Inches(9), Inches(0.8))
    subtitle_frame = subtitle_box.text_frame
    subtitle_frame.text = "Using Diagrams Library with Professional Icons"
    subtitle_para = subtitle_frame.paragraphs[0]
    subtitle_para.font.size = Pt(28)
    subtitle_para.font.color.rgb = DATABRICKS_DARK
    subtitle_para.alignment = PP_ALIGN.CENTER

    # Slide 2: Medallion Architecture (Simple)
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])

    # Title
    title_box2 = slide2.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame2 = title_box2.text_frame
    title_frame2.text = "Medallion Architecture - Simple View"
    title_para2 = title_frame2.paragraphs[0]
    title_para2.font.size = Pt(32)
    title_para2.font.bold = True
    title_para2.font.color.rgb = DATABRICKS_RED

    # Add diagram
    try:
        slide2.shapes.add_picture(
            "medallion_simple_poc.png",
            Inches(0.5),
            Inches(1.2),
            width=Inches(9)
        )
    except Exception as e:
        print(f"Warning: Could not add medallion_simple_poc.png: {e}")

    # Slide 3: Medallion Architecture (Detailed)
    slide3 = prs.slides.add_slide(prs.slide_layouts[6])

    # Title
    title_box3 = slide3.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame3 = title_box3.text_frame
    title_frame3.text = "Medallion Architecture - Detailed Components"
    title_para3 = title_frame3.paragraphs[0]
    title_para3.font.size = Pt(32)
    title_para3.font.bold = True
    title_para3.font.color.rgb = DATABRICKS_RED

    # Add diagram
    try:
        slide3.shapes.add_picture(
            "medallion_architecture_poc.png",
            Inches(0.5),
            Inches(1.2),
            width=Inches(9)
        )
    except Exception as e:
        print(f"Warning: Could not add medallion_architecture_poc.png: {e}")

    # Slide 4: Lakehouse Architecture
    slide4 = prs.slides.add_slide(prs.slide_layouts[6])

    # Title
    title_box4 = slide4.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(9), Inches(0.6))
    title_frame4 = title_box4.text_frame
    title_frame4.text = "Databricks Lakehouse Platform"
    title_para4 = title_frame4.paragraphs[0]
    title_para4.font.size = Pt(32)
    title_para4.font.bold = True
    title_para4.font.color.rgb = DATABRICKS_RED

    # Add diagram
    try:
        slide4.shapes.add_picture(
            "lakehouse_architecture_poc.png",
            Inches(0.5),
            Inches(1.2),
            width=Inches(9)
        )
    except Exception as e:
        print(f"Warning: Could not add lakehouse_architecture_poc.png: {e}")

    # Save presentation
    output_file = "enterprise_diagrams_poc.pptx"
    prs.save(output_file)
    print(f"✅ PowerPoint created: {output_file}")
    return output_file


if __name__ == "__main__":
    print("Creating PowerPoint with enterprise diagrams...")
    create_presentation_with_diagrams()
    print("\n✨ Success! Open 'enterprise_diagrams_poc.pptx' to see the results.")
