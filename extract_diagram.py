"""
Extract diagram images from PowerPoint for preview
"""
from pptx import Presentation
from PIL import Image
import io


def extract_diagrams(pptx_path):
    """Extract all diagram images from PowerPoint"""

    prs = Presentation(pptx_path)
    diagram_count = 0

    print("="*80)
    print("EXTRACTING DIAGRAMS FROM POWERPOINT")
    print("="*80)
    print()

    for slide_num, slide in enumerate(prs.slides, 1):
        print(f"Slide {slide_num}:")

        for shape_num, shape in enumerate(slide.shapes):
            # Check if it's a picture (diagram)
            if shape.shape_type == 13:  # Picture type
                diagram_count += 1
                print(f"  ✓ Found diagram #{diagram_count}")

                # Extract image
                try:
                    image = shape.image
                    image_bytes = image.blob

                    # Save to file
                    output_path = f"/tmp/diagram_{diagram_count}.png"
                    with open(output_path, 'wb') as f:
                        f.write(image_bytes)

                    # Get dimensions
                    img = Image.open(io.BytesIO(image_bytes))
                    width, height = img.size

                    print(f"    Saved to: {output_path}")
                    print(f"    Size: {width}x{height} pixels")
                    print(f"    Format: {img.format}")
                    print()

                except Exception as e:
                    print(f"    Error extracting: {e}")
                    print()

        if diagram_count == 0:
            print(f"  No diagrams on this slide")

    print("="*80)
    print(f"Total diagrams extracted: {diagram_count}")
    print("="*80)

    return diagram_count


if __name__ == "__main__":
    count = extract_diagrams("/tmp/genie_mcp_zerobus_demo.pptx")
