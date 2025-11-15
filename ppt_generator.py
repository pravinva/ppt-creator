"""
PowerPoint Generator with Databricks Branding
Creates professional presentations using python-pptx
"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from typing import List, Dict, Any
from diagram_generator import DiagramGenerator
import io


class PPTGenerator:
    """Generate PowerPoint presentations with Databricks branding"""

    # Databricks brand colors
    DATABRICKS_RED = RGBColor(255, 54, 33)
    DATABRICKS_DARK = RGBColor(27, 49, 57)
    DATABRICKS_WHITE = RGBColor(255, 255, 255)
    DATABRICKS_GRAY = RGBColor(102, 102, 102)

    def __init__(self):
        """Initialize PowerPoint generator"""
        self.prs = Presentation()
        self.prs.slide_width = Inches(10)
        self.prs.slide_height = Inches(7.5)
        self.diagram_gen = DiagramGenerator()

    def create_presentation(self, content: Dict[str, Any]) -> Presentation:
        """
        Create a complete presentation from structured content

        Args:
            content: Dictionary with title and slides

        Returns:
            Presentation object
        """
        # Create title slide
        self._add_title_slide(content.get('title', 'Databricks Solution Architecture'))

        # Create content slides
        for slide_data in content.get('slides', []):
            self._add_content_slide(
                title=slide_data.get('title', ''),
                content=slide_data.get('content', []),
                diagram_type=slide_data.get('diagram_type', 'none'),
                diagram_description=slide_data.get('diagram_description')
            )

        return self.prs

    def _add_title_slide(self, title: str):
        """Add title slide with Databricks branding"""
        slide_layout = self.prs.slide_layouts[6]  # Blank layout
        slide = self.prs.slides.add_slide(slide_layout)

        # Add Databricks-style background
        background = slide.shapes.add_shape(
            1,  # Rectangle
            0, 0,
            self.prs.slide_width,
            self.prs.slide_height
        )
        background.fill.solid()
        background.fill.fore_color.rgb = self.DATABRICKS_WHITE
        background.line.color.rgb = self.DATABRICKS_WHITE

        # Add red accent bar at top
        accent = slide.shapes.add_shape(
            1,  # Rectangle
            0, 0,
            self.prs.slide_width,
            Inches(0.3)
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = self.DATABRICKS_RED
        accent.line.color.rgb = self.DATABRICKS_RED

        # Add title
        title_box = slide.shapes.add_textbox(
            Inches(1),
            Inches(2.5),
            Inches(8),
            Inches(2)
        )
        title_frame = title_box.text_frame
        title_frame.text = title
        title_frame.paragraphs[0].alignment = PP_ALIGN.CENTER

        # Format title
        title_para = title_frame.paragraphs[0]
        title_para.font.size = Pt(54)
        title_para.font.bold = True
        title_para.font.color.rgb = self.DATABRICKS_DARK

        # Add subtitle
        subtitle_box = slide.shapes.add_textbox(
            Inches(1),
            Inches(5),
            Inches(8),
            Inches(0.8)
        )
        subtitle_frame = subtitle_box.text_frame
        subtitle_frame.text = "Solutions Architecture"
        subtitle_frame.paragraphs[0].alignment = PP_ALIGN.CENTER
        subtitle_para = subtitle_frame.paragraphs[0]
        subtitle_para.font.size = Pt(28)
        subtitle_para.font.color.rgb = self.DATABRICKS_GRAY

    def _add_content_slide(
        self,
        title: str,
        content: List[str],
        diagram_type: str = 'none',
        diagram_description: str = None
    ):
        """Add a content slide with optional diagram"""
        slide_layout = self.prs.slide_layouts[6]  # Blank layout
        slide = self.prs.slides.add_slide(slide_layout)

        # Add background
        background = slide.shapes.add_shape(
            1,  # Rectangle
            0, 0,
            self.prs.slide_width,
            self.prs.slide_height
        )
        background.fill.solid()
        background.fill.fore_color.rgb = self.DATABRICKS_WHITE
        background.line.color.rgb = self.DATABRICKS_WHITE

        # Add red accent bar at top
        accent = slide.shapes.add_shape(
            1,  # Rectangle
            0, 0,
            self.prs.slide_width,
            Inches(0.15)
        )
        accent.fill.solid()
        accent.fill.fore_color.rgb = self.DATABRICKS_RED
        accent.line.color.rgb = self.DATABRICKS_RED

        # Add title
        title_box = slide.shapes.add_textbox(
            Inches(0.5),
            Inches(0.4),
            Inches(9),
            Inches(0.8)
        )
        title_frame = title_box.text_frame
        title_frame.text = title
        title_para = title_frame.paragraphs[0]
        title_para.font.size = Pt(36)
        title_para.font.bold = True
        title_para.font.color.rgb = self.DATABRICKS_DARK

        # Determine layout based on diagram
        if diagram_type and diagram_type != 'none':
            # Two-column layout: content left, diagram right
            self._add_bullet_points(slide, content, Inches(0.5), Inches(1.5), Inches(4.5), Inches(5.5))
            self._add_diagram(slide, diagram_type, diagram_description, Inches(5.2), Inches(1.5), Inches(4.3), Inches(5.5))
        else:
            # Full-width content
            self._add_bullet_points(slide, content, Inches(0.5), Inches(1.5), Inches(9), Inches(5.5))

    def _add_bullet_points(self, slide, content: List[str], left, top, width, height):
        """Add bullet points to slide"""
        content_box = slide.shapes.add_textbox(left, top, width, height)
        text_frame = content_box.text_frame
        text_frame.word_wrap = True

        for i, item in enumerate(content):
            if i == 0:
                p = text_frame.paragraphs[0]
            else:
                p = text_frame.add_paragraph()

            p.text = item
            p.level = 0
            p.font.size = Pt(20)
            p.font.color.rgb = self.DATABRICKS_DARK
            p.space_before = Pt(12)

    def _add_diagram(self, slide, diagram_type: str, description: str, left, top, width, height):
        """Add diagram to slide"""
        if diagram_type == 'medallion':
            diagram_img = self.diagram_gen.generate_medallion_architecture(description)
        elif diagram_type == 'architecture':
            diagram_img = self.diagram_gen.generate_architecture_diagram(description)
        else:
            return  # No diagram

        # Convert image to bytes
        img_stream = self.diagram_gen.get_image_bytes(diagram_img)

        # Add image to slide
        pic = slide.shapes.add_picture(img_stream, left, top, width=width)

    def save_presentation(self, filepath: str):
        """Save presentation to file"""
        self.prs.save(filepath)

    def get_presentation_bytes(self) -> bytes:
        """Get presentation as bytes for download"""
        ppt_stream = io.BytesIO()
        self.prs.save(ppt_stream)
        ppt_stream.seek(0)
        return ppt_stream.getvalue()
